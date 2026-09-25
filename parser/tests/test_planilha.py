"""Nova importação Excel — planilha própria do cliente com as colunas escolhidas na mão."""

from __future__ import annotations

import io
import json
from datetime import date, datetime
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.parsers.planilha import (
    PlanilhaInvalidaError,
    _centavos_texto,
    _data_texto,
    letra_coluna,
    ler_planilha,
    parse_planilha,
)
from app.schemas import ExcelMapeamento

FIX = Path(__file__).parent / "fixtures"
client = TestClient(app)

# A = Data, B = Histórico, C = Fornecedor/Cliente, D = Valor, E = Saldo
MAPA = dict(aba=0, data=0, valor=3, historico=[1, 2])


def controle_xlsx() -> bytes:
    """Controle financeiro de cliente, com as esquisitices de planilha feita à mão.

    tests/fixtures/controle_cliente.xls é ESTA planilha salva como .xls pelo Excel.
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Agosto"
    ws["A1"] = "CONTROLE FINANCEIRO - AGOSTO/2026"
    ws.merge_cells("A1:E1")
    ws.append([])  # linha 2 em branco
    ws.append(["Data", "Histórico", "Fornecedor/Cliente", "Valor", "Saldo"])  # 3
    ws.append(["01/08/2026", "Saldo anterior", None, 1500, 1500])  # 4: data em texto
    ws.append([datetime(2026, 8, 3), "Venda balcão", "Cliente A", 350.5, 1850.5])  # 5
    ws.append([datetime(2026, 8, 5), "Pix fornecedor", "Fornecedor B", -1234.56, 615.94])  # 6
    ws.append([None, "Tarifa", None, "R$ -12,90", 603.04])  # 7: data mesclada com a de cima
    ws.merge_cells("A6:A7")
    ws.append([datetime(2026, 8, 10), "Estorno", None, 0, 603.04])  # 8: valor zero
    ws.append([datetime(2026, 8, 12), "Aluguel", "Imobiliária C", "(2.000,00)", -1396.96])  # 9
    serial = (date(2026, 8, 15) - date(1899, 12, 30)).days
    ws.append([serial, "Depósito", None, 800, -596.96])  # 10: data como nº de série, sem formato
    ws.append([datetime(2026, 8, 20), None, None, -50, -646.96])  # 11: sem histórico
    ws.append([datetime(2026, 8, 25), "Parcela 1/2", None, -300, -946.96])  # 12
    ws.append([datetime(2026, 8, 26), "Parcela 2/2", None, None, -946.96])  # 13: valor mesclado
    ws.merge_cells("D12:D13")
    ws.append([None, "TOTAL", None, 1234, None])  # 14: total sem data
    ws.append([])  # 15
    ws.append(["Observação: conferir notas"])  # 16

    resumo = wb.create_sheet("Resumo")
    resumo.append(["Data", "Valor", "Descrição"])
    resumo.append([datetime(2026, 8, 31), 99.9, "Outro"])
    resumo.sheet_state = "hidden"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cel(g, linha: int, col: int):
    row = next(r for r in g.linhas if r.n == linha)
    return row.c[col] if col < len(row.c) else None


def _totais(r):
    e = [t for t in r.transactions if t.direction == "entrada"]
    s = [t for t in r.transactions if t.direction == "saida"]
    return len(e), sum(t.amount_cents for t in e), len(s), sum(t.amount_cents for t in s)


# --------------------------------------------------------------------------- #
# grade
# --------------------------------------------------------------------------- #
def test_grade_mostra_a_aba_com_os_numeros_de_linha_do_excel():
    g = ler_planilha(controle_xlsx())
    assert g.formato == "xlsx"
    assert [(a.nome, a.oculta) for a in g.abas] == [("Agosto", False), ("Resumo", True)]
    assert g.aba == 0
    assert g.colunas == 5
    # linhas 2 e 15 em branco não vêm
    assert [r.n for r in g.linhas] == [1, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 16]
    assert g.total_linhas == 14 and not g.truncado


def test_grade_le_cada_celula_como_texto_data_e_valor():
    g = ler_planilha(controle_xlsx())
    assert _cel(g, 3, 0).model_dump() == {"t": "Data", "d": None, "v": None}
    assert _cel(g, 4, 0).model_dump() == {"t": "01/08/2026", "d": "2026-08-01", "v": None}
    assert _cel(g, 5, 0).model_dump() == {"t": "03/08/2026", "d": "2026-08-03", "v": None}
    assert _cel(g, 5, 3).model_dump() == {"t": "350,5", "d": None, "v": 35050}
    assert _cel(g, 6, 3).v == -123456
    assert _cel(g, 7, 3).model_dump() == {"t": "R$ -12,90", "d": None, "v": -1290}
    assert _cel(g, 9, 3).v == -200000
    assert _cel(g, 10, 0).d == "2026-08-15"  # nº de série
    assert _cel(g, 5, 1).t == "Venda balcão"
    assert _cel(g, 11, 1) is None


def test_mesclagem_vertical_repete_a_data_mas_nao_o_valor():
    g = ler_planilha(controle_xlsx())
    assert _cel(g, 7, 0).d == "2026-08-05"  # A6:A7
    assert _cel(g, 12, 3).v == -30000
    assert _cel(g, 13, 3).t == "-300" and _cel(g, 13, 3).v is None  # D12:D13
    # horizontal (título A1:E1) não se espalha pelas colunas
    assert _cel(g, 1, 0).t.startswith("CONTROLE") and _cel(g, 1, 1) is None


def test_sugere_as_colunas_pelo_cabecalho():
    s = ler_planilha(controle_xlsx()).sugestao
    assert (s.data, s.valor, s.historico) == (0, 3, [1])


def test_nao_sugere_valor_quando_ha_coluna_de_entrada_e_de_saida():
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["Data", "Descrição", "Valor entrada", "Valor saída"])
    wb.active.append([datetime(2026, 8, 1), "Venda", 10, None])
    buf = io.BytesIO()
    wb.save(buf)
    s = ler_planilha(buf.getvalue()).sugestao
    assert (s.data, s.valor, s.historico) == (0, None, [1])


def test_com_varias_colunas_de_valor_sugere_a_que_se_chama_so_valor():
    # layout de sistema financeiro (ex.: "extrato financeiro" exportado de ERP)
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(
        ["Data movimento", "Descrição", "Valor original (R$)", "Valor (R$)", "Saldo conta (R$)", "Valor na Categoria 1"]
    )
    wb.active.append([datetime(2026, 8, 1), "Venda", 10, 10, 110, 10])
    buf = io.BytesIO()
    wb.save(buf)
    s = ler_planilha(buf.getvalue()).sugestao
    assert (s.data, s.valor, s.historico) == (0, 3, [1])

    # "Valor previsto" x "Valor pago": nenhuma é só "Valor" — o operador escolhe
    wb = Workbook()
    wb.active.append(["Vencimento", "Descrição", "Valor Previsto (R$)", "Data Liquidação", "Valor Pago (R$)"])
    buf = io.BytesIO()
    wb.save(buf)
    s = ler_planilha(buf.getvalue()).sugestao
    assert (s.data, s.valor, s.historico) == (3, None, [1])


def test_escolhe_a_aba():
    g = ler_planilha(controle_xlsx(), aba=1)
    assert g.aba == 1 and g.abas[1].nome == "Resumo"
    assert _cel(g, 2, 1).v == 9990
    # aba que não existe na leitura da grade: volta pra aba ativa
    assert ler_planilha(controle_xlsx(), aba=7).aba == 0


# --------------------------------------------------------------------------- #
# lançamentos pelas colunas escolhidas
# --------------------------------------------------------------------------- #
def test_sinal_do_valor_decide_entrada_ou_saida():
    r = parse_planilha(controle_xlsx(), ExcelMapeamento(**MAPA))
    assert r.format == "xlsx"
    assert [(t.date, t.direction, t.amount_cents, t.description, t.raw["linha"]) for t in r.transactions] == [
        ("2026-08-01", "entrada", 150000, "Saldo anterior", 4),
        ("2026-08-03", "entrada", 35050, "Venda balcão - Cliente A", 5),
        ("2026-08-05", "saida", 123456, "Pix fornecedor - Fornecedor B", 6),
        ("2026-08-05", "saida", 1290, "Tarifa", 7),
        ("2026-08-12", "saida", 200000, "Aluguel - Imobiliária C", 9),
        ("2026-08-15", "entrada", 80000, "Depósito", 10),
        ("2026-08-20", "saida", 5000, "", 11),
        ("2026-08-25", "saida", 30000, "Parcela 1/2", 12),
    ]
    assert r.transactions[0].raw["aba"] == "Agosto"
    assert (r.period_start, r.period_end) == ("2026-08-01", "2026-08-25")
    assert r.warnings == [
        "4 linha(s) sem data na coluna A ficaram de fora (cabeçalho, títulos...): 1, 3, 14, 16",
        "1 linha(s) sem valor na coluna D ficaram de fora: 13",
        "1 linha(s) com valor zero ficaram de fora: 8",
        "1 lançamento(s) sem histórico — linha(s) 11",
    ]


def test_linhas_tiradas_pelo_operador_ficam_de_fora():
    r = parse_planilha(controle_xlsx(), ExcelMapeamento(**{**MAPA, "excluir": [4, 12]}))
    assert _totais(r) == (2, 35050 + 80000, 4, 123456 + 1290 + 200000 + 5000)
    assert "2 linha(s) tiradas da importação por você: 4, 12" in r.warnings
    # tirar linha que já ficaria de fora (cabeçalho) não muda nada
    r2 = parse_planilha(controle_xlsx(), ExcelMapeamento(**{**MAPA, "excluir": [3]}))
    assert len(r2.transactions) == 8


def test_historico_junta_as_colunas_na_ordem_da_planilha():
    r = parse_planilha(controle_xlsx(), ExcelMapeamento(**{**MAPA, "historico": [2, 1]}))
    assert r.transactions[1].description == "Venda balcão - Cliente A"


def test_aba_inexistente_na_importacao_e_erro():
    with pytest.raises(PlanilhaInvalidaError, match="não tem a aba 3"):
        parse_planilha(controle_xlsx(), ExcelMapeamento(**{**MAPA, "aba": 2}))


def test_coluna_que_nao_existe_e_erro():
    with pytest.raises(PlanilhaInvalidaError, match="coluna H não existe"):
        parse_planilha(controle_xlsx(), ExcelMapeamento(**{**MAPA, "historico": [7]}))


def test_mesma_coluna_com_duas_funcoes_e_erro():
    with pytest.raises(PlanilhaInvalidaError, match="uma função"):
        parse_planilha(controle_xlsx(), ExcelMapeamento(**{**MAPA, "valor": 0}))
    with pytest.raises(PlanilhaInvalidaError, match="uma função"):
        parse_planilha(controle_xlsx(), ExcelMapeamento(**{**MAPA, "historico": [1, 3]}))


def test_arquivo_que_nao_e_planilha():
    with pytest.raises(PlanilhaInvalidaError, match="não é uma planilha Excel"):
        ler_planilha(b"Data;Valor\n01/08/2026;10,00\n")


# --------------------------------------------------------------------------- #
# .xls (Excel 97-2003): mesma planilha, mesmo resultado
# --------------------------------------------------------------------------- #
def test_xls_le_igual_ao_xlsx():
    xls = (FIX / "controle_cliente.xls").read_bytes()
    g_xls, g_xlsx = ler_planilha(xls), ler_planilha(controle_xlsx())
    assert g_xls.formato == "xls"
    assert [(a.nome, a.oculta) for a in g_xls.abas] == [("Agosto", False), ("Resumo", True)]
    assert g_xls.model_dump(exclude={"formato"}) == g_xlsx.model_dump(exclude={"formato"})

    r_xls = parse_planilha(xls, ExcelMapeamento(**MAPA))
    r_xlsx = parse_planilha(controle_xlsx(), ExcelMapeamento(**MAPA))
    assert r_xls.model_dump(exclude={"format"}) == r_xlsx.model_dump(exclude={"format"})


# --------------------------------------------------------------------------- #
# célula digitada como texto
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize(
    "texto, centavos",
    [
        ("1.234,56", 123456),
        ("R$ 1.234,56", 123456),
        ("R$ -1.234,56", -123456),
        ("-R$ 1.234,56", -123456),
        ("(1.234,56)", -123456),
        ("1.234,56-", -123456),
        ("150,00 D", -15000),
        ("150,00 C", 15000),
        ("+10,00", 1000),
        ("1234.56", 123456),
        ("1,234.56", 123456),
        ("1.500", 150000),  # só ponto + 3 dígitos = milhar (regra do sistema)
        ("10.5", 1050),
        ("1.234.567", 123456700),
        (",50", 50),
        ("0", 0),
        ("01/08/2026", None),
        ("PIX 50", None),
        ("12.34,56", None),
        ("", None),
        ("-", None),
    ],
)
def test_valor_em_texto(texto, centavos):
    assert _centavos_texto(texto) == centavos


@pytest.mark.parametrize(
    "texto, iso",
    [
        ("01/08/2026", "2026-08-01"),
        ("1/8/2026", "2026-08-01"),
        ("01/08/26", "2026-08-01"),
        ("01.08.2026", "2026-08-01"),
        ("01-08-2026", "2026-08-01"),
        ("2026-08-01", "2026-08-01"),
        ("17/07/2026 17:03:07", "2026-07-17"),
        ("31/02/2026", None),
        ("Data", None),
        ("1/12", None),
    ],
)
def test_data_em_texto(texto, iso):
    d = _data_texto(texto)
    assert (d.isoformat() if d else None) == iso


def test_letra_da_coluna():
    assert [letra_coluna(i) for i in (0, 25, 26, 27, 51, 52)] == ["A", "Z", "AA", "AB", "AZ", "BA"]


# --------------------------------------------------------------------------- #
# endpoints
# --------------------------------------------------------------------------- #
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_endpoint_planilha_omite_o_que_nao_e_data_nem_valor():
    r = client.post("/excel/planilha", files={"file": ("controle.xlsx", controle_xlsx(), XLSX_MIME)})
    assert r.status_code == 200
    body = r.json()
    assert body["aba"] == 0 and body["colunas"] == 5
    linha3 = next(l for l in body["linhas"] if l["n"] == 3)
    assert linha3["c"][0] == {"t": "Data"}
    linha5 = next(l for l in body["linhas"] if l["n"] == 5)
    assert linha5["c"][0] == {"t": "03/08/2026", "d": "2026-08-03"}
    assert linha5["c"][3] == {"t": "350,5", "v": 35050}
    linha11 = next(l for l in body["linhas"] if l["n"] == 11)
    assert linha11["c"][1] is None  # célula vazia continua na lista (posição = coluna)
    assert body["sugestao"] == {"data": 0, "valor": 3, "historico": [1]}


def test_endpoint_planilha_escolhe_a_aba():
    r = client.post(
        "/excel/planilha",
        files={"file": ("controle.xlsx", controle_xlsx(), XLSX_MIME)},
        data={"aba": "1"},
    )
    assert r.status_code == 200 and r.json()["aba"] == 1


def test_endpoint_parse_excel():
    r = client.post(
        "/parse/excel",
        files={"file": ("controle.xlsx", controle_xlsx(), XLSX_MIME)},
        data={"mapeamento": json.dumps({**MAPA, "excluir": [4]})},
    )
    assert r.status_code == 200
    body = r.json()
    assert body["format"] == "xlsx" and len(body["transactions"]) == 7


def test_endpoint_parse_excel_mapeamento_invalido_da_422():
    r = client.post(
        "/parse/excel",
        files={"file": ("controle.xlsx", controle_xlsx(), XLSX_MIME)},
        data={"mapeamento": json.dumps({"aba": 0, "data": 0, "valor": 3, "historico": []})},
    )
    assert r.status_code == 422 and r.json()["code"] == "planilha"

    r = client.post(
        "/parse/excel",
        files={"file": ("controle.xlsx", controle_xlsx(), XLSX_MIME)},
        data={"mapeamento": json.dumps({**MAPA, "historico": [9]})},
    )
    assert r.status_code == 422
    assert r.json() == {"error": 'a coluna J não existe na aba "Agosto"', "code": "planilha"}


def test_endpoint_arquivo_que_nao_e_planilha_da_422():
    r = client.post("/excel/planilha", files={"file": ("x.xls", b"Data;Valor\n", "application/vnd.ms-excel")})
    assert r.status_code == 422 and r.json()["code"] == "planilha"
