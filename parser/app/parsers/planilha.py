"""Planilha Excel própria do cliente (controle interno) — "Nova importação Excel".

Não tenta adivinhar o layout como os leitores de extrato: o operador diz qual
coluna é Data, Valor e Histórico. Duas etapas:

  1. `ler_planilha` devolve a aba como grade. Cada célula vem com o texto e já
     lida como data (`d`) e como valor (`v`, centavos com sinal) quando der.
     A tela mostra a grade, o operador escolhe as colunas e vê na hora o que
     vai entrar.
  2. `parse_planilha` aplica a escolha e devolve o ParseResult de sempre.
     Valor negativo = saída, positivo = entrada.

As duas passam pela mesma leitura de célula (`_celula`), então a prévia da
tela — que só aplica a escolha sobre `d`/`v` — bate com o que é gravado.
"""

from __future__ import annotations

import io
import math
import re
from datetime import date, datetime, time, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

from ..schemas import (
    ExcelMapeamento,
    NormalizedTransaction,
    ParseResult,
    PlanilhaAba,
    PlanilhaCelula,
    PlanilhaLinha,
    PlanilhaResult,
    PlanilhaSugestao,
)
from .excel import EncryptedFileError, _is_encrypted_ole2
from .tabular import _norm

# A prévia vai até aqui: acima do limite de 10.000 lançamentos por importação,
# com folga pra cabeçalho, títulos e totais.
MAX_LINHAS = 12_000
MAX_COLUNAS = 60


class PlanilhaInvalidaError(Exception):
    """Não é planilha Excel legível, ou a escolha de colunas não serve pra ela."""


class _Mesclada:
    """Célula coberta por uma mesclagem vertical (ex.: a mesma data pra várias
    linhas): mostra o conteúdo de cima, mas não conta como valor — senão um
    valor mesclado viraria um lançamento por linha."""

    __slots__ = ("x",)

    def __init__(self, x: object):
        self.x = x


# --------------------------------------------------------------------------- #
# abrir a planilha → matriz de valores crus (uma lista por linha, a partir da 1)
# --------------------------------------------------------------------------- #
def _escolher_aba(pedida: int | None, n: int, padrao: int, estrita: bool) -> int:
    if pedida is None:
        return padrao
    if 0 <= pedida < n:
        return pedida
    if estrita:
        raise PlanilhaInvalidaError(f"a planilha não tem a aba {pedida + 1}")
    return padrao


def _abrir(content: bytes, aba: int | None, estrita: bool = False):
    """→ (formato, abas, índice da aba lida, matriz, epoch 1904?, mesclas)."""
    if content[:4] == b"\xd0\xcf\x11\xe0":
        return _abrir_xls(content, aba, estrita)
    if content[:2] == b"PK":
        return _abrir_xlsx(content, aba, estrita)
    raise PlanilhaInvalidaError(
        "o arquivo não é uma planilha Excel de verdade (pode ser um CSV ou HTML com extensão "
        ".xls) — abra no Excel e salve como .xlsx"
    )


def _abrir_xlsx(content: bytes, aba: int | None, estrita: bool):
    from openpyxl import load_workbook
    from openpyxl.utils.datetime import CALENDAR_MAC_1904

    try:
        # read_only=False: o modo streaming trunca algumas planilhas (ver excel.py)
        wb = load_workbook(io.BytesIO(content), read_only=False, data_only=True, keep_links=False)
    except Exception as e:  # noqa: BLE001
        raise PlanilhaInvalidaError(f"não consegui abrir a planilha ({e})") from e
    try:
        folhas = wb.worksheets
        if not folhas:
            raise PlanilhaInvalidaError("a planilha não tem nenhuma aba")
        abas = [PlanilhaAba(nome=ws.title, oculta=ws.sheet_state != "visible") for ws in folhas]
        ativa = folhas.index(wb.active) if wb.active in folhas else 0
        idx = _escolher_aba(aba, len(folhas), ativa, estrita)
        ws = folhas[idx]
        # min_row/min_col=1: sem isso a openpyxl começa na 1ª linha preenchida
        # e o número da linha deixa de bater com o do Excel
        matriz = [
            list(r)
            for r in ws.iter_rows(
                min_row=1, min_col=1, max_row=ws.max_row, max_col=ws.max_column, values_only=True
            )
        ]
        mesclas = [
            (m.min_row - 1, m.max_row - 1, m.min_col - 1, m.max_col - 1) for m in ws.merged_cells.ranges
        ]
        epoch1904 = wb.epoch == CALENDAR_MAC_1904
    finally:
        wb.close()
    return "xlsx", abas, idx, matriz, epoch1904, mesclas


def _abrir_xls(content: bytes, aba: int | None, estrita: bool):
    if _is_encrypted_ole2(content):
        raise EncryptedFileError(
            "planilha protegida por senha — tire a senha no Excel e suba de novo"
        )
    import xlrd

    try:
        try:
            # formatting_info traz as células mescladas
            book = xlrd.open_workbook(file_contents=content, formatting_info=True)
        except Exception:  # noqa: BLE001
            book = xlrd.open_workbook(file_contents=content)
    except xlrd.XLRDError as e:
        raise EncryptedFileError(
            f"não consegui abrir a planilha ({e}). Se ela tem senha, tire a senha no Excel."
        ) from e

    folhas = book.sheets()
    if not folhas:
        raise PlanilhaInvalidaError("a planilha não tem nenhuma aba")
    abas = [PlanilhaAba(nome=s.name, oculta=s.visibility != 0) for s in folhas]
    padrao = next((i for i, s in enumerate(folhas) if s.visibility == 0), 0)
    idx = _escolher_aba(aba, len(folhas), padrao, estrita)
    sh = folhas[idx]

    vazias = (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR)
    matriz: list[list[object]] = []
    for r in range(sh.nrows):
        linha: list[object] = []
        for c in range(sh.ncols):
            cel = sh.cell(r, c)
            if cel.ctype in vazias:
                linha.append(None)
            elif cel.ctype == xlrd.XL_CELL_DATE:
                try:
                    linha.append(xlrd.xldate_as_datetime(cel.value, book.datemode))
                except Exception:  # noqa: BLE001
                    linha.append(cel.value)
            elif cel.ctype == xlrd.XL_CELL_BOOLEAN:
                linha.append(bool(cel.value))
            else:
                linha.append(cel.value)
        matriz.append(linha)
    mesclas = [(rlo, rhi - 1, clo, chi - 1) for rlo, rhi, clo, chi in getattr(sh, "merged_cells", [])]
    return "xls", abas, idx, matriz, book.datemode == 1, mesclas


def _aplicar_mesclas(matriz: list[list[object]], mesclas) -> None:
    """Mesclagem vertical: repete o conteúdo de cima nas linhas cobertas (só na
    1ª coluna da mesclagem — como o Excel mostra). Horizontal fica como está."""
    for r0, r1, c0, _c1 in mesclas:
        if r1 <= r0 or r0 >= len(matriz) or c0 >= len(matriz[r0]):
            continue
        x = matriz[r0][c0]
        if x is None or (isinstance(x, str) and not x.strip()):
            continue
        for r in range(r0 + 1, min(r1, len(matriz) - 1) + 1):
            if c0 < len(matriz[r]):
                matriz[r][c0] = _Mesclada(x)


# --------------------------------------------------------------------------- #
# leitura de uma célula: texto + data + valor
# --------------------------------------------------------------------------- #
_ESPACOS = re.compile(r"\s+")
_DATA_BR = re.compile(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{4}|\d{2})(?:[ T]+\d{1,2}:\d{2}(?::\d{2})?)?$")
_DATA_ISO = re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T]+\d{1,2}:\d{2}(?::\d{2}(?:\.\d+)?)?)?$")
_SO_NUMERO = re.compile(r"^[\d.,]*\d$")


def _data_texto(s: str) -> date | None:
    """'01/08/2026', '1/8/26', '01.08.2026', '01-08-2026', '2026-08-01' (com ou sem hora)."""
    m = _DATA_BR.match(s)
    if m:
        dia, mes, ano = int(m[1]), int(m[2]), int(m[3])
        if len(m[3]) == 2:
            ano += 2000 if ano < 70 else 1900
    else:
        m = _DATA_ISO.match(s)
        if not m:
            return None
        ano, mes, dia = int(m[1]), int(m[2]), int(m[3])
    try:
        return date(ano, mes, dia)
    except ValueError:
        return None


def _data_serial(x: float, epoch1904: bool) -> date | None:
    """Número de série do Excel numa célula sem formato de data (45870 = 01/08/2025)."""
    if not math.isfinite(x) or x < 1:
        return None
    base = date(1904, 1, 1) if epoch1904 else date(1899, 12, 30)
    try:
        d = base + timedelta(days=int(x))
    except OverflowError:
        return None
    return d if 1990 <= d.year <= 2100 else None


def _grupos_de_milhar(partes: list[str]) -> bool:
    return 1 <= len(partes[0]) <= 3 and all(len(p) == 3 for p in partes[1:])


def _centavos_texto(s: str) -> int | None:
    """Valor digitado como texto → centavos com sinal.

    '1.234,56', 'R$ -1.234,56', '(1.234,56)', '1.234,56-', '150,00 D' (débito),
    '150,00 C', '1234.56', '1,234.56'. Só ponto: 3 dígitos depois dele é milhar
    ('1.500' = 1500), senão decimal ('10.5') — a mesma regra do resto do sistema.
    """
    t = _ESPACOS.sub("", s.upper().replace("R$", "").replace("−", "-"))
    neg = False
    if t.startswith("(") and t.endswith(")"):
        neg, t = True, t[1:-1]
    if len(t) > 1 and t[-1] in "DC" and t[-2].isdigit():
        neg, t = neg or t[-1] == "D", t[:-1]
    if t[:1] in ("-", "+"):
        neg, t = neg or t[0] == "-", t[1:]
    elif t[-1:] in ("-", "+"):
        neg, t = neg or t[-1] == "-", t[:-1]
    if not _SO_NUMERO.match(t):
        return None

    if "," in t and "." in t:
        dec = "," if t.rfind(",") > t.rfind(".") else "."
        inteiro, _, frac = t.rpartition(dec)
        if dec in inteiro or not _grupos_de_milhar(inteiro.split("." if dec == "," else ",")):
            return None
        inteiro = inteiro.replace("." if dec == "," else ",", "")
    elif "," in t or "." in t:
        sep = "," if "," in t else "."
        partes = t.split(sep)
        if len(partes) > 2:
            if not _grupos_de_milhar(partes):
                return None
            inteiro, frac = "".join(partes), ""
        elif sep == "." and len(partes[1]) == 3:
            inteiro, frac = partes[0] + partes[1], ""
        else:
            inteiro, frac = partes
    else:
        inteiro, frac = t, ""

    try:
        reais = Decimal(f"{inteiro or '0'}.{frac or '0'}")
    except InvalidOperation:
        return None
    cents = int((reais * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    return -cents if neg else cents


def _centavos_numero(x: float | int) -> int | None:
    if isinstance(x, float) and not math.isfinite(x):
        return None
    # repr: o float 1234.56 vira "1234.56", não 1234.5599999...
    return int((Decimal(repr(x)) * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _texto_numero(x: float | int) -> str:
    if isinstance(x, int) or x.is_integer():
        return str(int(x))
    s = format(x, ".15g")
    return (repr(x) if "e" in s else s).replace(".", ",")


def _dma(d: date) -> str:
    return f"{d.day:02d}/{d.month:02d}/{d.year:04d}"


def _cel(t: str, d: date | None = None, v: int | None = None) -> PlanilhaCelula:
    # model_construct: sem validação — uma planilha grande tem dezenas de milhares de células
    return PlanilhaCelula.model_construct(t=t, d=d.isoformat() if d else None, v=v)


def _celula(x: object, epoch1904: bool) -> PlanilhaCelula | None:
    """Valor cru da planilha → célula (texto, data, valor). None = vazia."""
    if x is None:
        return None
    if isinstance(x, _Mesclada):
        c = _celula(x.x, epoch1904)
        return c and PlanilhaCelula.model_construct(t=c.t, d=c.d, v=None)
    if isinstance(x, bool):
        return _cel("VERDADEIRO" if x else "FALSO")
    if isinstance(x, datetime):
        hora = f" {x.hour:02d}:{x.minute:02d}" if x.time() != time(0) else ""
        return _cel(_dma(x) + hora, d=x.date())
    if isinstance(x, date):
        return _cel(_dma(x), d=x)
    if isinstance(x, time):
        return _cel(f"{x.hour:02d}:{x.minute:02d}")
    if isinstance(x, (int, float)):
        if isinstance(x, float) and not math.isfinite(x):
            return _cel(str(x))
        return _cel(_texto_numero(x), d=_data_serial(float(x), epoch1904), v=_centavos_numero(x))
    s = _ESPACOS.sub(" ", str(x)).strip()
    if not s:
        return None
    d = _data_texto(s)
    return _cel(s, d=d, v=None if d else _centavos_texto(s))


def _vazia(bruta: list[object]) -> bool:
    return all(x is None or (isinstance(x, str) and not x.strip()) for x in bruta[:MAX_COLUNAS])


def letra_coluna(i: int) -> str:
    """0 → A, 25 → Z, 26 → AA."""
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


# --------------------------------------------------------------------------- #
# 1. grade pra tela escolher as colunas
# --------------------------------------------------------------------------- #
def _titulo_data(t: str) -> bool:
    return len(t) <= 40 and (t in ("dt", "date") or t.startswith(("data", "dt ", "dt.")))


def _titulo_valor(t: str) -> bool:
    return len(t) <= 40 and "saldo" not in t and (
        "valor" in t or t in ("vlr", "vlr.", "vl", "montante", "amount", "quantia", "value", "r$")
    )


def _titulo_historico(t: str) -> bool:
    return len(t) <= 40 and (
        any(k in t for k in ("histor", "descri", "memo", "detalhe")) or t in ("lancamento", "lancamentos")
    )


def _sugerir(linhas: list[PlanilhaLinha]) -> PlanilhaSugestao | None:
    """Colunas pelo cabeçalho (linha com título de data E de valor). Com várias
    colunas de valor, só sugere a que se chama exatamente "Valor" (ex.: "Valor
    (R$)" ao lado de "Valor original", "Valor na categoria") — "Valor entrada"/
    "Valor saída" ou "Valor previsto"/"Valor pago" fica pro operador decidir."""
    for linha in linhas[:30]:
        titulos = [_norm(c.t) if c else "" for c in linha.c]
        data = next((i for i, t in enumerate(titulos) if _titulo_data(t)), None)
        valores = [i for i, t in enumerate(titulos) if i != data and _titulo_valor(t)]
        if data is None or not valores:
            continue
        exatos = [i for i in valores if titulos[i] in ("valor", "valor (r$)", "valor(r$)", "valor r$")]
        valor = valores[0] if len(valores) == 1 else exatos[0] if len(exatos) == 1 else None
        historico = [i for i, t in enumerate(titulos) if i not in (data, valor) and _titulo_historico(t)]
        return PlanilhaSugestao(data=data, valor=valor, historico=historico)
    return None


def ler_planilha(content: bytes, aba: int | None = None) -> PlanilhaResult:
    formato, abas, idx, matriz, epoch1904, mesclas = _abrir(content, aba)
    _aplicar_mesclas(matriz, mesclas)

    linhas: list[PlanilhaLinha] = []
    total = 0
    colunas = 0
    for i, bruta in enumerate(matriz):
        if _vazia(bruta):
            continue
        total += 1
        if len(linhas) >= MAX_LINHAS:
            continue
        cels = [_celula(x, epoch1904) for x in bruta[:MAX_COLUNAS]]
        while cels and cels[-1] is None:
            cels.pop()
        colunas = max(colunas, len(cels))
        linhas.append(PlanilhaLinha.model_construct(n=i + 1, c=cels))

    return PlanilhaResult(
        formato=formato,
        abas=abas,
        aba=idx,
        colunas=colunas,
        linhas=linhas,
        total_linhas=total,
        truncado=total > len(linhas),
        sugestao=_sugerir(linhas),
    )


# --------------------------------------------------------------------------- #
# 2. aplica a escolha do operador → lançamentos
# --------------------------------------------------------------------------- #
def _lista(ns: list[int], limite: int = 15) -> str:
    s = ", ".join(str(n) for n in ns[:limite])
    return s + (f" e mais {len(ns) - limite}" if len(ns) > limite else "")


def parse_planilha(content: bytes, mapa: ExcelMapeamento) -> ParseResult:
    if mapa.data == mapa.valor or mapa.data in mapa.historico or mapa.valor in mapa.historico:
        raise PlanilhaInvalidaError("cada coluna só pode ter uma função (Data, Valor ou Histórico)")

    formato, abas, idx, matriz, epoch1904, mesclas = _abrir(content, mapa.aba, estrita=True)
    largura = min(max((len(r) for r in matriz), default=0), MAX_COLUNAS)
    maior = max(mapa.data, mapa.valor, *mapa.historico)
    if maior >= largura:
        raise PlanilhaInvalidaError(f"a coluna {letra_coluna(maior)} não existe na aba \"{abas[idx].nome}\"")
    _aplicar_mesclas(matriz, mesclas)

    historico = sorted(set(mapa.historico))
    excluir = set(mapa.excluir)
    txns: list[NormalizedTransaction] = []
    datas: list[date] = []
    sem_data: list[int] = []
    sem_valor: list[int] = []
    zerados: list[int] = []
    tiradas: list[int] = []
    sem_historico: list[int] = []

    def celula(bruta: list[object], c: int) -> PlanilhaCelula | None:
        return _celula(bruta[c], epoch1904) if c < len(bruta) else None

    # mesma ordem de checagem da prévia da tela (features/import/excel/planilha.ts)
    for i, bruta in enumerate(matriz):
        if _vazia(bruta):
            continue
        n = i + 1
        cd, cv = celula(bruta, mapa.data), celula(bruta, mapa.valor)
        d = cd.d if cd else None
        v = cv.v if cv else None
        if not d:
            sem_data.append(n)
            continue
        if v is None:
            sem_valor.append(n)
            continue
        if v == 0:
            zerados.append(n)
            continue
        if n in excluir:
            tiradas.append(n)
            continue
        partes = [c.t for c in (celula(bruta, h) for h in historico) if c and c.t]
        descricao = " - ".join(partes)
        if not descricao:
            sem_historico.append(n)
        txns.append(
            NormalizedTransaction(
                date=d,
                description=descricao,
                amount_cents=abs(v),
                direction="saida" if v < 0 else "entrada",
                raw={"linha": n, "aba": abas[idx].nome},
            )
        )
        datas.append(date.fromisoformat(d))

    L = letra_coluna
    warnings: list[str] = []
    if sem_data:
        warnings.append(
            f"{len(sem_data)} linha(s) sem data na coluna {L(mapa.data)} ficaram de fora "
            f"(cabeçalho, títulos...): {_lista(sem_data)}"
        )
    if sem_valor:
        warnings.append(
            f"{len(sem_valor)} linha(s) sem valor na coluna {L(mapa.valor)} ficaram de fora: {_lista(sem_valor)}"
        )
    if zerados:
        warnings.append(f"{len(zerados)} linha(s) com valor zero ficaram de fora: {_lista(zerados)}")
    if tiradas:
        warnings.append(f"{len(tiradas)} linha(s) tiradas da importação por você: {_lista(tiradas)}")
    if sem_historico:
        warnings.append(f"{len(sem_historico)} lançamento(s) sem histórico — linha(s) {_lista(sem_historico)}")

    return ParseResult(
        format=formato,
        period_start=min(datas).isoformat() if datas else None,
        period_end=max(datas).isoformat() if datas else None,
        transactions=txns,
        warnings=warnings,
    )
